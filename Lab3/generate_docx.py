#!/usr/bin/env python3
"""Generate a .docx lab report matching the style of ZPO_LAB3_Lesniewski_Wiktor.pdf."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_run_font(run, name="Times New Roman", size=12, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), name)


def add_paragraph(doc, text="", bold=False, size=12, align=None, font_name="Times New Roman", italic=False, space_after=6):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    if text:
        run = p.add_run(text)
        set_run_font(run, font_name, size, bold, italic)
    return p


def add_code_block(doc, code, size=9):
    for line in code.split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.left_indent = Inches(0.2)
        shading = OxmlElement('w:shd')
        shading.set(qn('w:val'), 'clear')
        shading.set(qn('w:fill'), 'F0F0F0')
        shading.set(qn('w:color'), 'auto')
        pPr = p._element.get_or_add_pPr()
        pPr.append(shading)
        run = p.add_run(line if line else " ")
        set_run_font(run, "Consolas", size)
    add_paragraph(doc, "", size=6)


def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_run_font(run, "Times New Roman", 12)
    return p


def add_heading_custom(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, "Times New Roman", 12, True)
    return h


def add_hr(doc):
    p = doc.add_paragraph()
    pPr = p._element.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '000000')
    pBdr.append(bottom)
    pPr.append(pBdr)


def build_docx():
    doc = Document()

    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.paragraph_format.space_after = Pt(6)

    # ===================== TITLE PAGE =====================
    add_paragraph(doc, "", size=12)
    add_paragraph(doc, "Politechnika Bydgoska im. J. J. Sniadeckich", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "Wydzial Telekomunikacji, Informatyki i Elektrotechniki", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "", size=6)
    add_hr(doc)
    add_paragraph(doc, "", size=6)

    rows = [
        ("Przedmiot", "Zaawansowane Programowanie Obiektowe"),
        ("Prowadzacy", "dr inz. Damian Szczegielniak"),
        ("Temat", "Implementacja aplikacji bazodanowej przy uzyciu JDBC + JavaFX"),
        ("Student", "Wiktor Lesniewski"),
        ("Nr cw.", "3"),
        ("Data wykonania", "maj 2026"),
        ("Ocena", ""),
        ("Data oddania spr.", ""),
    ]
    for label, value in rows:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(label + ":  ")
        set_run_font(run, "Times New Roman", 12, True)
        run = p.add_run(value)
        set_run_font(run, "Times New Roman", 12)

    add_hr(doc)
    add_paragraph(doc, "", size=6)

    # ===================== SECTION: OPIS PROBLEMU =====================
    add_heading_custom(doc, "Opis problemu do rozwiazania:", level=2)

    add_paragraph(doc,
        "Celem zadania bylo stworzenie w pelni funkcjonalnej aplikacji okienkowej w jezyku Java, "
        "laczacej sie z relacyjna baza danych HSQLDB za pomoca technologii JDBC. "
        "Aplikacja umozliwia wykonywanie operacji CRUD (Create, Read, Update, Delete) "
        "na dwoch powiazanych ze soba tabelach: 'projekt' oraz 'zadanie'. "
        "Dodatkowo zaimplementowano mechanizm logowania (SLF4J + Logback), "
        "zarzadzanie polaczeniami poprzez pule HikariCP oraz graficzny interfejs "
        "uzytkownika w technologii JavaFX z wykorzystaniem plikow FXML."
    )

    add_paragraph(doc, "W projekcie zaimplementowano nastepujace komponenty:")

    bullets = [
        "Model danych (Projekt, Zadanie) - klasy Java odwzorowujace tabele bazy danych.",
        "DataSource (Singleton) - pula polaczen oparta na HikariCP.",
        "DbInitializer - tworzy strukture tabel, indeksow i kluczy obcych.",
        "ProjektDAO, ZadanieDAO - interfejsy i implementacje z PreparedStatement.",
        "ProjectFrame.fxml, ZadanieFrame.fxml - interfejs graficzny JavaFX.",
        "ProjectController, ZadanieController - obsluga zdarzen GUI i logika biznesowa.",
        "Logback + SLF4J - mechanizm rejestracji z archiwizacja logow.",
    ]
    for b in bullets:
        add_bullet(doc, b)

    # ===================== STRUKTURA =====================
    doc.add_page_break()
    add_heading_custom(doc, "Struktura projektu", level=2)
    add_code_block(doc, """\
project-jfx-client/
+-- build.gradle
+-- settings.gradle
+-- src/main/java/
|   +-- module-info.java
|   +-- com/project/
|       +-- app/ProjectClientApplication.java
|       +-- controller/ProjectController.java
|       +-- controller/ZadanieController.java
|       +-- dao/ProjektDAO.java
|       +-- dao/ProjektDAOImpl.java
|       +-- dao/ZadanieDAO.java
|       +-- dao/ZadanieDAOImpl.java
|       +-- datasource/DataSource.java
|       +-- datasource/DbInitializer.java
|       +-- model/Projekt.java
|       +-- model/Zadanie.java
+-- src/main/resources/
    +-- css/application.css
    +-- fxml/ProjectFrame.fxml
    +-- fxml/ZadanieFrame.fxml
    +-- logback.xml""")

    # ===================== BUILD GRADLE =====================
    add_heading_custom(doc, "Konfiguracja projektu (build.gradle)", level=2)
    add_paragraph(doc,
        "Plik build.gradle okresla zaleznosci: HSQLDB 2.7.4, HikariCP 6.2.1, "
        "Logback 1.5.17, SLF4J 2.0.17, JavaFX 23.0.2. Zgodnosc z Java 21.")
    add_code_block(doc, """\
plugins {
    id 'application'  id 'java'  id 'eclipse'
    id 'org.openjfx.javafxplugin' version '0.1.0'
    id 'org.beryx.jlink' version '3.1.1'
}
group = 'com.project'  version = '1.0'
repositories { mavenCentral() }
dependencies {
    implementation 'org.hsqldb:hsqldb:2.7.4'
    implementation 'com.zaxxer:HikariCP:6.2.1'
    implementation 'ch.qos.logback:logback-classic:1.5.17'
    implementation 'ch.qos.logback:logback-core:1.5.17'
    implementation 'org.slf4j:slf4j-api:2.0.17'
}
java { sourceCompatibility = '21'  modularity.inferModulePath = true }
javafx { version = '23.0.2'
    modules = ['javafx.controls', 'javafx.fxml', 'javafx.base', 'javafx.graphics'] }
application {
    mainModule = 'project.jfx.client'
    mainClass = 'com.project.app.ProjectClientApplication'
}""")

    # ===================== LOGBACK =====================
    add_heading_custom(doc, "Mechanizm rejestracji (Logback)", level=2)
    add_paragraph(doc,
        "Plik logback.xml definiuje appenders STDOUT (konsola) i FILE (plik z archiwizacja). "
        "Logger com.project zapisuje do obu, root tylko na konsole. "
        "Archiwizacja dzienna, 30 dni historii, max 30 MB.")
    add_code_block(doc, """\
<configuration debug="true">
    <property name="LOG_FILE" value="project-jfx-client" />
    <property name="LOG_DIR" value="logs" />
    <appender name="STDOUT" class="ch.qos.logback.core.ConsoleAppender">
        <encoder><pattern>%d{HH:mm:ss.SSS} [%thread] %-5level %logger{36}.%M\\(%line\\) - %msg%n</pattern></encoder>
    </appender>
    <appender name="FILE" class="ch.qos.logback.core.rolling.RollingFileAppender">
        <file>${LOG_DIR}/${LOG_FILE}.log</file>
        <rollingPolicy class="ch.qos.logback.core.rolling.TimeBasedRollingPolicy">
            <fileNamePattern>${LOG_ARCHIVE}/%d{yyyy-MM-dd}${LOG_FILE}.log.zip</fileNamePattern>
            <maxHistory>30</maxHistory>
            <totalSizeCap>30MB</totalSizeCap>
        </rollingPolicy>
    </appender>
    <logger name="com.project" level="INFO" additivity="false">
        <appender-ref ref="STDOUT" /><appender-ref ref="FILE" />
    </logger>
    <root level="INFO"><appender-ref ref="STDOUT" /></root>
</configuration>""")

    # ===================== MODEL =====================
    doc.add_page_break()
    add_heading_custom(doc, "Model danych - klasy Projekt i Zadanie", level=2)
    add_paragraph(doc,
        "Klasy Java odwzorowujace tabele bazy. Projekt: projektId, nazwa, opis, "
        "dataCzasUtworzenia (LocalDateTime), dataOddania (LocalDate). "
        "Zadanie: zadanieId, nazwa, opis, kolejnosc, dataCzasUtworzenia, projektId (klucz obcy).")
    add_code_block(doc, """\
public class Projekt {
    private Integer projektId;
    private String nazwa;
    private String opis;
    private LocalDateTime dataCzasUtworzenia;
    private LocalDate dataOddania;
    public Projekt() {}
    public Projekt(Integer projektId, String nazwa, String opis,
                   LocalDateTime dataCzasUtworzenia, LocalDate dataOddania) { ... }
    public Projekt(String nazwa, String opis,
                   LocalDateTime dataCzasUtworzenia, LocalDate dataOddania) { ... }
}""")
    add_code_block(doc, """\
public class Zadanie {
    private Integer zadanieId;
    private String nazwa;
    private String opis;
    private Integer kolejnosc;
    private LocalDateTime dataCzasUtworzenia;
    private Integer projektId;
}""")

    # ===================== DATASOURCE =====================
    add_heading_custom(doc, "Pula polaczen - DataSource (Singleton)", level=2)
    add_paragraph(doc,
        "Klasa DataSource implementuje Singleton dla HikariDataSource. "
        "URL do lokalnej bazy HSQLDB z parametrami write_delay=false i sql.syntax_pgs=true. "
        "Rozmiar puli = 1.")
    add_code_block(doc, """\
public class DataSource {
    private final static String DB_DIR = "db";
    private final static String DB_NAME = "projekty";
    private final static String DB_USERNAME = "admin";
    private final static String DB_USER_PASSWORD = "admin";
    private final static String HSQL_ADDITIONAL_PARAMS =
        ";hsqldb.write_delay=false;sql.syntax_pgs=true";
    private final static String DB_URL =
        String.format("jdbc:hsqldb:file:%s/%s%s", DB_DIR, DB_NAME, HSQL_ADDITIONAL_PARAMS);
    private final static HikariDataSource ds;
    static {
        HikariConfig config = new HikariConfig();
        config.setJdbcUrl(DB_URL);
        config.setUsername(DB_USERNAME);
        config.setPassword(DB_USER_PASSWORD);
        config.setMaximumPoolSize(1);
        ds = new HikariDataSource(config);
    }
    private DataSource() {}
    public static Connection getConnection() throws SQLException {
        return ds.getConnection();
    }
}""")

    # ===================== DB INITIALIZER =====================
    add_heading_custom(doc, "Inicjalizacja bazy danych - DbInitializer", level=2)
    add_paragraph(doc,
        "Tworzy strukture bazy podczas uruchamiania. Wykorzystuje transakcje: "
        "auto-commit = false, execute, commit. Blad -> rollback. "
        "Tabele z IF NOT EXISTS, klucz obcy z ON DELETE CASCADE.")
    add_code_block(doc, """\
CREATE TABLE IF NOT EXISTS projekt (
    projekt_id SERIAL,  nazwa VARCHAR(50) NOT NULL,
    opis VARCHAR(1000),  dataczas_utworzenia TIMESTAMP DEFAULT now(),
    data_oddania DATE,  CONSTRAINT projekt_pk PRIMARY KEY (projekt_id)
);
CREATE TABLE IF NOT EXISTS zadanie (
    zadanie_id SERIAL,  nazwa VARCHAR(50) NOT NULL,
    opis VARCHAR(1000),  kolejnosc INTEGER,
    dataczas_utworzenia TIMESTAMP DEFAULT now(),
    projekt_id INTEGER NOT NULL,
    CONSTRAINT zadanie_pk PRIMARY KEY (zadanie_id)
);
CREATE INDEX IF NOT EXISTS projekt_nazwa_idx ON projekt(nazwa);
CREATE INDEX IF NOT EXISTS zadanie_nazwa_idx ON zadanie(nazwa);
ALTER TABLE zadanie ADD CONSTRAINT IF NOT EXISTS zadanie_projekt_fk
    FOREIGN KEY (projekt_id) REFERENCES projekt (projekt_id) ON DELETE CASCADE;
ALTER TABLE zadanie ADD CONSTRAINT IF NOT EXISTS unique_kolejnosc UNIQUE (kolejnosc, projekt_id);""")

    # ===================== DAO =====================
    doc.add_page_break()
    add_heading_custom(doc, "Warstwa DAO - interfejsy i implementacje", level=2)
    add_paragraph(doc,
        "Interfejs ProjektDAO definiuje metody CRUD z paginacja i filtrowaniem. "
        "Implementacja uzywa PreparedStatement (SQL Injection prevention).")
    add_code_block(doc, """\
public interface ProjektDAO {
    Projekt getProjekt(Integer projektId);
    void setProjekt(Projekt projekt);
    void deleteProjekt(Integer projektId);
    List<Projekt> getProjekty(Integer offset, Integer limit);
    List<Projekt> getProjektyWhereNazwaLike(String nazwa, Integer offset, Integer limit);
    List<Projekt> getProjektyWhereDataOddaniaIs(LocalDate d, Integer offset, Integer limit);
    int getRowsNumber();
    int getRowsNumberWhereNazwaLike(String nazwa);
    int getRowsNumberWhereDataOddaniaIs(LocalDate d);
}""")

    add_paragraph(doc, "Metoda setProjekt (INSERT/UPDATE z pobraniem klucza):", bold=True)
    add_code_block(doc, """\
@Override
public void setProjekt(Projekt p) {
    boolean insert = p.getProjektId() == null;
    String sql = insert ?
        "INSERT INTO projekt(nazwa,opis,dataczas_utworzenia,data_oddania) VALUES(?,?,?,?)"
        : "UPDATE projekt SET nazwa=?,opis=?,dataczas_utworzenia=?,data_oddania=?"
          + " WHERE projekt_id=?";
    try (Connection c = DataSource.getConnection();
         PreparedStatement ps = c.prepareStatement(sql, Statement.RETURN_GENERATED_KEYS)) {
        ps.setString(1, p.getNazwa());
        ps.setString(2, p.getOpis());
        if (p.getDataCzasUtworzenia() == null)
            p.setDataCzasUtworzenia(LocalDateTime.now());
        ps.setObject(3, p.getDataCzasUtworzenia());
        ps.setObject(4, p.getDataOddania());
        if (!insert) ps.setInt(5, p.getProjektId());
        ps.executeUpdate();
        if (insert && ps.getUpdateCount() > 0) {
            try (ResultSet rs = ps.getGeneratedKeys()) {
                if (rs.next()) p.setProjektId(rs.getInt(1));
            }
        }
    } catch (Exception e) { throw new RuntimeException(e); }
}""")

    add_paragraph(doc, "Metoda getProjekty z paginacja (LIMIT/OFFSET):", bold=True)
    add_code_block(doc, """\
@Override
public List<Projekt> getProjekty(Integer offset, Integer limit) {
    List<Projekt> list = new ArrayList<>();
    try (Connection c = DataSource.getConnection();
         PreparedStatement ps = c.prepareStatement(
             "SELECT * FROM projekt ORDER BY dataczas_utworzenia DESC LIMIT ? OFFSET ?")) {
        ps.setInt(1, limit);   ps.setInt(2, offset);
        ResultSet rs = ps.executeQuery();
        while (rs.next()) {
            Projekt p = new Projekt();
            p.setProjektId(rs.getInt(1));
            p.setNazwa(rs.getString(2));
            p.setOpis(rs.getString(3));
            p.setDataCzasUtworzenia(rs.getObject(4, LocalDateTime.class));
            p.setDataOddania(rs.getObject(5, LocalDate.class));
            list.add(p);
        }
    } catch (Exception e) { throw new RuntimeException(e); }
    return list;
}""")

    # ===================== INTERFEJS GRAFICZNY =====================
    doc.add_page_break()
    add_heading_custom(doc, "Interfejs graficzny (JavaFX FXML)", level=2)
    add_paragraph(doc,
        "Glowne okno (ProjectFrame.fxml): pole wyszukiwania, tabela z projektami, "
        "przyciski nawigacji, ChoiceBox rozmiaru strony. "
        "Okno zadan (ZadanieFrame.fxml): lista zadan dla projektu.")
    add_code_block(doc, """\
<!-- ProjectFrame.fxml -->
<BorderPane prefHeight="600" prefWidth="900"
    xmlns:fx="http://javafx.com/fxml/1"
    fx:controller="com.project.controller.ProjectController">
    <top>
        <HBox alignment="CENTER_LEFT" spacing="10">
            <Label text="Wyszukaj:" />
            <TextField fx:id="txtSzukaj" promptText="Nazwa, data lub ID..." />
            <Button onAction="#onActionBtnSzukaj" text="Szukaj" />
            <Region HBox.hgrow="ALWAYS" />
            <Button onAction="#onActionBtnDodaj" text="Dodaj projekt" />
        </HBox>
    </top>
    <center>
        <TableView fx:id="tblProjekt">
            <columns>
                <TableColumn fx:id="colId" text="ID" />
                <TableColumn fx:id="colNazwa" text="Nazwa" />
                <TableColumn fx:id="colOpis" text="Opis" />
                <TableColumn fx:id="colDataCzasUtworzenia" text="Utworzono" />
                <TableColumn fx:id="colDataOddania" text="Termin" />
            </columns>
        </TableView>
    </center>
    <bottom>
        <HBox alignment="CENTER" spacing="10">
            <Button fx:id="btnPierwsza" onAction="#onActionBtnPierwsza" text="|\u25c4" />
            <Button fx:id="btnWstecz" onAction="#onActionBtnWstecz" text="\u25c4" />
            <Label text="Rozmiar strony:" />
            <ChoiceBox fx:id="cbPageSizes" />
            <Button fx:id="btnDalej" onAction="#onActionBtnDalej" text="\u25ba" />
            <Button fx:id="btnOstatnia" onAction="#onActionBtnOstatnia" text="\u25ba|" />
        </HBox>
    </bottom>
</BorderPane>""")

    add_code_block(doc, """\
<!-- ZadanieFrame.fxml -->
<BorderPane prefHeight="500" prefWidth="800"
    xmlns:fx="http://javafx.com/fxml/1"
    fx:controller="com.project.controller.ZadanieController">
    <top>
        <HBox alignment="CENTER_LEFT" spacing="10">
            <Label fx:id="lblTytul" text="Zadania dla projektu" />
            <Region HBox.hgrow="ALWAYS" />
            <Button onAction="#onActionBtnDodaj" text="Dodaj zadanie" />
        </HBox>
    </top>
    <center>
        <TableView fx:id="tblZadanie">
            <columns>
                <TableColumn fx:id="colId" text="ID" />
                <TableColumn fx:id="colNazwa" text="Nazwa" />
                <TableColumn fx:id="colOpis" text="Opis" />
                <TableColumn fx:id="colKolejnosc" text="Kolejnosc" />
                <TableColumn fx:id="colDataCzasUtworzenia" text="Utworzono" />
            </columns>
        </TableView>
    </center>
    <bottom>
        <Button fx:id="btnPowrot" onAction="#onActionBtnPowrot" text="Powrot" />
    </bottom>
</BorderPane>""")

    # ===================== KONTROLER =====================
    doc.add_page_break()
    add_heading_custom(doc, "Warstwa kontrolera - ProjectController", level=2)
    add_paragraph(doc,
        "Klasa zarzadza GUI i obsluga zdarzen. Uzywa ExecutorService do asynchronicznego "
        "ladowania danych, ObservableList do automatycznej synchronizacji tabeli.")

    add_paragraph(doc, "Inicjalizacja:", bold=True)
    add_code_block(doc, """\
@FXML public void initialize() {
    search4 = "";  pageNo = 0;  pageSize = 10;
    cbPageSizes.getItems().addAll(5, 10, 20, 50, 100);
    cbPageSizes.setValue(pageSize);
    cbPageSizes.setOnAction(e -> { pageSize = cbPageSizes.getValue(); pageNo = 0;
        wykonawca.execute(() -> loadPage(search4, pageNo, pageSize)); });
    colId.setCellValueFactory(new PropertyValueFactory<>("projektId"));
    colNazwa.setCellValueFactory(new PropertyValueFactory<>("nazwa"));
    colOpis.setCellValueFactory(new PropertyValueFactory<>("opis"));
    colDataCzasUtworzenia.setCellValueFactory(new PropertyValueFactory<>("dataCzasUtworzenia"));
    colDataOddania.setCellValueFactory(new PropertyValueFactory<>("dataOddania"));
    colDataCzasUtworzenia.setCellFactory(column -> new TableCell<>() {
        @Override protected void updateItem(LocalDateTime item, boolean empty) {
            super.updateItem(item, empty);
            setText(empty || item == null ? null : dateTimeFormatter.format(item));
        }
    });
    projekty = FXCollections.observableArrayList();
    tblProjekt.setItems(projekty);
    wykonawca.execute(() -> loadPage(search4, pageNo, pageSize));
}""")

    add_paragraph(doc, "Wczytywanie danych z paginacja i wyszukiwaniem (wyrazenia regularne):", bold=True)
    add_code_block(doc, """\
private void loadPage(String search4, Integer pageNo, Integer pageSize) {
    try {
        List<Projekt> lista = new ArrayList<>();
        if (search4 != null && !search4.isEmpty()) {
            if (search4.matches("[0-9]+")) {
                Projekt p = projektDAO.getProjekt(Integer.parseInt(search4));
                if (p != null) lista.add(p);
            } else if (search4.matches("^\\\\d{4}-\\\\d{2}-\\\\d{2}$")) {
                lista.addAll(projektDAO.getProjektyWhereDataOddaniaIs(
                    LocalDate.parse(search4), pageNo * pageSize, pageSize));
            } else {
                lista.addAll(projektDAO.getProjektyWhereNazwaLike(
                    search4, pageNo * pageSize, pageSize));
            }
        } else {
            lista.addAll(projektDAO.getProjekty(pageNo * pageSize, pageSize));
        }
        Platform.runLater(() -> { projekty.clear(); projekty.addAll(lista); });
    } catch (Exception e) { logger.error("Blad ladowania danych", e); }
}""")

    add_paragraph(doc, "Kolumna z przyciskami (Edycja, Usun, Zadania):", bold=True)
    add_code_block(doc, """\
TableColumn<Projekt, Void> colEdit = new TableColumn<>("Edycja");
colEdit.setCellFactory(column -> new TableCell<>() {
    private final GridPane pane;
    {
        Button btnTask = new Button("Zadania");
        Button btnEdit = new Button("Edycja");
        Button btnRemove = new Button("Usun");
        btnTask.setOnAction(e -> openZadanieFrame(getCurrentProjekt()));
        btnEdit.setOnAction(e -> edytujProjekt(getCurrentProjekt()));
        btnRemove.setOnAction(e -> usunProjekt(getCurrentProjekt()));
        pane = new GridPane();
        pane.add(btnTask, 0, 0);  pane.add(btnEdit, 0, 1);  pane.add(btnRemove, 0, 2);
    }
    private Projekt getCurrentProjekt() { return getTableView().getItems().get(getIndex()); }
    @Override protected void updateItem(Void item, boolean empty) { setGraphic(empty ? null : pane); }
});
tblProjekt.getColumns().add(colEdit);""")

    add_paragraph(doc, "Dialog dodawania/edycji projektu:", bold=True)
    add_code_block(doc, """\
private void edytujProjekt(Projekt projekt) {
    Dialog<Projekt> dialog = new Dialog<>();
    dialog.setTitle("Edycja");
    TextField txtNazwa = new TextField(projekt.getNazwa());
    TextArea txtOpis = new TextArea(projekt.getOpis());
    DatePicker datePicker = new DatePicker(projekt.getDataOddania());
    GridPane grid = new GridPane();
    grid.add(new Label("Nazwa:"), 0, 0);  grid.add(txtNazwa, 1, 0);
    grid.add(new Label("Opis:"), 0, 1);   grid.add(txtOpis, 1, 1);
    grid.add(new Label("Data oddania:"), 0, 2);  grid.add(datePicker, 1, 2);
    dialog.getDialogPane().setContent(grid);
    ButtonType ok = new ButtonType("Zapisz", ButtonBar.ButtonData.OK_DONE);
    dialog.getDialogPane().getButtonTypes().addAll(ok, ButtonType.CANCEL);
    dialog.setResultConverter(btn -> {
        if (btn == ok) {
            projekt.setNazwa(txtNazwa.getText());
            projekt.setOpis(txtOpis.getText());
            projekt.setDataOddania(datePicker.getValue());
            return projekt;
        } return null;
    });
    Optional<Projekt> result = dialog.showAndWait();
    result.ifPresent(p -> wykonawca.execute(() -> {
        projektDAO.setProjekt(p);
        Platform.runLater(() -> {
            if (!projekty.contains(p)) projekty.add(0, p);
            tblProjekt.refresh();
        });
    }));
}""")

    add_paragraph(doc, "Usuwanie projektu z potwierdzeniem i shutdown:", bold=True)
    add_code_block(doc, """\
private void usunProjekt(Projekt projekt) {
    Alert alert = new Alert(AlertType.CONFIRMATION);
    alert.setHeaderText("Usunac projekt?");
    alert.setContentText(projekt.getNazwa());
    Optional<ButtonType> result = alert.showAndWait();
    if (result.isPresent() && result.get() == ButtonType.OK) {
        wykonawca.execute(() -> {
            projektDAO.deleteProjekt(projekt.getProjektId());
            Platform.runLater(() -> projekty.remove(projekt));
        });
    }
}
public void shutdown() {
    if (wykonawca != null) {
        wykonawca.shutdown();
        try { if (!wykonawca.awaitTermination(5, TimeUnit.SECONDS))
                wykonawca.shutdownNow();
        } catch (InterruptedException e) { wykonawca.shutdownNow(); }
    }
}""")

    # ===================== ZADANIE CONTROLLER =====================
    doc.add_page_break()
    add_heading_custom(doc, "Kontroler zadan - ZadanieController", level=2)
    add_paragraph(doc,
        "Klasa obsluguje okno zadan. Konstruktor przyjmuje Projekt, ZadanieDAO, ExecutorService. "
        "loadZadania() pobiera zadania z bazy. Pelny CRUD. Przycisk Powrot zamyka okno.")
    add_code_block(doc, """\
public class ZadanieController {
    private final Projekt projekt;
    private final ZadanieDAO zadanieDAO;
    private final ExecutorService wykonawca;
    @FXML private Button btnPowrot;
    @FXML private Label lblTytul;
    @FXML private TableView<Zadanie> tblZadanie;
    public ZadanieController(Projekt p, ZadanieDAO z, ExecutorService w) {
        this.projekt = p;  this.zadanieDAO = z;  this.wykonawca = w;
    }
    @FXML public void initialize() {
        lblTytul.setText("Zadania dla projektu: " + projekt.getNazwa());
        wykonawca.execute(this::loadZadania);
    }
    private void loadZadania() {
        List<Zadanie> lista = zadanieDAO.getZadania(projekt.getProjektId());
        Platform.runLater(() -> { zadania.clear(); zadania.addAll(lista); });
    }
    @FXML private void onActionBtnPowrot(ActionEvent event) {
        Stage stage = (Stage) btnPowrot.getScene().getWindow();
        stage.fireEvent(new WindowEvent(stage, WindowEvent.WINDOW_CLOSE_REQUEST));
    }
}""")

    # ===================== PROJECT CLIENT APPLICATION =====================
    add_heading_custom(doc, "Klasa uruchomieniowa - ProjectClientApplication", level=2)
    add_paragraph(doc,
        "Rozszerza Application. main(): DbInitializer.init() + launch(). "
        "start(): laduje FXML, tworzy DAO, przekazuje do kontrolera przez "
        "setControllerFactory, rejestruje shutdown przy zamknieciu.")
    add_code_block(doc, """\
public class ProjectClientApplication extends Application {
    @Override
    public void start(Stage stage) throws Exception {
        FXMLLoader loader = new FXMLLoader(
            getClass().getResource("/fxml/ProjectFrame.fxml"));
        ProjektDAO dao = new ProjektDAOImpl();
        loader.setControllerFactory(c -> new ProjectController(dao));
        Scene scene = new Scene(loader.load());
        scene.getStylesheets().add(
            getClass().getResource("/css/application.css").toExternalForm());
        ProjectController controller = loader.getController();
        stage.setOnCloseRequest(e -> { controller.shutdown(); Platform.exit(); });
        stage.setTitle("Projekty");  stage.setScene(scene);  stage.show();
    }
    public static void main(String[] args) {
        DbInitializer.init();
        launch(args);
    }
}""")

    # ===================== MODULE-INFO =====================
    add_heading_custom(doc, "Modul Java - module-info.java", level=2)
    add_code_block(doc, """\
module project.jfx.client {
    exports com.project.datasource;
    exports com.project.dao;
    exports com.project.model;
    exports com.project.app;
    exports com.project.controller;
    requires javafx.base;
    requires javafx.fxml;
    requires javafx.controls;
    requires transitive javafx.graphics;
    requires com.zaxxer.hikari;
    requires transitive java.sql;
    requires org.hsqldb;
    requires org.slf4j;
    requires ch.qos.logback.classic;
    requires ch.qos.logback.core;
    opens com.project.app to javafx.graphics, javafx.fxml, javafx.base, javafx.controls;
    opens com.project.model to javafx.graphics, javafx.fxml, javafx.base, javafx.controls;
    opens com.project.controller to javafx.graphics, javafx.fxml, javafx.base, javafx.controls;
}""")

    # ===================== ODPOWIEDZI =====================
    doc.add_page_break()
    add_heading_custom(doc, "Opis problemu do rozwiazania - odpowiedzi na pytania", level=2)

    qa = [
        ("Jakie biblioteki zostaly uzyte?",
         "HSQLDB 2.7.4 (baza danych), HikariCP 6.2.1 (pula polaczen), "
         "Logback 1.5.17 + SLF4J 2.0.17 (logowanie), JavaFX 23.0.2 (GUI)."),
        ("Do czego sluzy plik build.gradle?",
         "Plik konfiguracyjny Gradle. Okresla zaleznosci, wersje, wtyczki, "
         "kodowanie, klase glowna i modul. Umożliwia zbudowanie i uruchomienie "
         "projektu jednym poleceniem."),
        ("Co to jest HikariCP i dlaczego go uzywamy?",
         "Lekka biblioteka do zarzadzania pula polaczen. Zamiast tworzyc "
         "nowe polaczenie za kazdym razem, pula przechowuje gotowe polaczenia, "
         "co poprawia wydajnosc. W aplikacji max. rozmiar puli = 1."),
        ("Czym jest wzorzec Singleton w klasie DataSource?",
         "Singleton ogranicza utworzenie tylko jednej instancji. "
         "Prywatny konstruktor uniemozliwia tworzenie instancji z zewnatrz, "
         "statyczny blok inicjalizujacy tworzy pojedynczy HikariDataSource."),
        ("Dlaczego uzywamy PreparedStatement zamiast Statement?",
         "PreparedStatement zapobiega SQL Injection poprzez automatyczne "
         "escapowanie parametrow. Pozwala na prekompilacje zapytania, "
         "co poprawia wydajnosc przy wielokrotnym wykonywaniu."),
        ("Jak dziala transakcja w DbInitializer?",
         "Wylacza auto-commit, wykonuje wszystkie zapytania, commit. "
         "Przy bledzie - rollback. Operacje sa atomowe: wszystkie albo zadna."),
        ("Jak dziala stronicowanie w aplikacji?",
         "Przez OFFSET i LIMIT w SQL. pageNo * pageSize = OFFSET, "
         "pageSize = LIMIT. Przyciski modyfikuja pageNo. ChoiceBox "
         "zmienia pageSize (5, 10, 20, 50, 100)."),
        ("Jak dziala wyszukiwanie?",
         "Wyrazenia regularne: cyfry -> szukaj po ID; "
         "RRRR-MM-DD -> szukaj po dacie oddania; "
         "reszta -> szukaj po nazwie (LIKE '%fraza%')."),
        ("Jak dziala kolumna 'Edycja' z przyciskami?",
         "Dodatkowa TableColumn z CellFactory. Kazda komorka "
         "zawiera GridPane z 3 przyciskami: Zadania, Edycja, Usun."),
        ("Do czego sluzy ExecutorService?",
         "Zarzadza pula watkow. Jednowatkowa pula wykonuje zadania w tle, "
         "nie blokujac watku JavaFX Application Thread. "
         "shutdown() zapewnia grzeczne zakonczenie watkow."),
    ]
    for q, a in qa:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(q)
        set_run_font(run, "Times New Roman", 12, True)
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(6)
        run = p2.add_run(a)
        set_run_font(run, "Times New Roman", 12)

    # ===================== WNIOSKI =====================
    doc.add_page_break()
    add_heading_custom(doc, "Wnioski", level=2)
    add_paragraph(doc,
        "Realizacja zadania pozwolila na praktyczne zapoznanie sie z budowa aplikacji "
        "bazodanowej w jezyku Java z wykorzystaniem nowoczesnych bibliotek i wzorcow "
        "projektowych:")

    wnioski = [
        "Stosowanie JDBC z PreparedStatement zapewnia bezpieczenstwo i wydajnosc zapytan.",
        "HikariCP jako pula polaczen eliminuje narzut zwiazany z tworzeniem nowych polaczen.",
        "Wzorzec DAO hermetyzuje logike dostepu do danych, ulatwiajac testowanie i modyfikacje.",
        "JavaFX z FXML i Scene Builder umozliwia szybkie projektowanie interfejsu uzytkownika.",
        "ExecutorService pozwala na asynchroniczne ladowanie danych bez blokowania GUI.",
        "ObservableList zapewnia automatyczna synchronizacje danych z widokiem tabeli.",
        "Mechanizm logowania (Logback) ulatwia diagnozowanie bledow w aplikacji.",
        "Wyrazenia regularne pozwalaja na elastyczne wyszukiwanie (ID, data, nazwa).",
        "Transakcje bazodanowe gwarantuja spojnosc danych przy tworzeniu struktury BD.",
        "System budowania Gradle upraszcza zarzadzanie zaleznosciami i budowanie projektu.",
    ]
    for w in wnioski:
        add_bullet(doc, w)

    add_paragraph(doc, "")
    add_paragraph(doc,
        "Repozytorium: https://github.com/Gaz3ll/Zaawansowane_Programowanie_Obiektowe",
        italic=True)

    out = "D:/eclipse-2024-09/workspace/project-jfx-client/sprawozdanie.docx"
    doc.save(out)
    print(f"DOCX saved: {out}")


if __name__ == "__main__":
    build_docx()
